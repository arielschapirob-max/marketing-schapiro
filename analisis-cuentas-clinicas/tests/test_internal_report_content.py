"""Verifica que el informe interno SÍ contenga el detalle y la terminología
central exigida (a diferencia de la propuesta comercial, que debe omitirlos).
Ver tests/test_commercial_proposal_no_leak.py para el caso contrario.
"""

import types

from docx import Document as DocxDocument

from app.db.models import Hallazgo, ItemCuenta
from app.reports.internal_report_docx import generar_informe_interno_docx


def _caso_falso():
    return types.SimpleNamespace(
        nombre_cliente="Cliente de Prueba",
        numero_cuenta="CTA-2026-0099",
        isapre="Consalud",
        prestador="Hospital Clínico Ejemplo",
        rut_cliente="22.222.222-2",
        aprobado_por_abogado=True,
        fecha_aprobacion=None,
    )


def _item():
    return ItemCuenta(
        id=1,
        codigo_prestacion="220305",
        descripcion="Stent coronario liberador de fármaco",
        valor_cobrado=3200000.0,
        valor_bonificado=1800000.0,
        copago=1400000.0,
    )


def _hallazgo(item):
    return Hallazgo(
        id=1,
        item_id=item.id,
        tipo="dispositivo_medico",
        monto_discutible=3200000.0,
        prioridad="alta",
        estado="pendiente",
        explicacion_interna="Existen antecedentes preliminares que ameritan evaluación.",
        documentos_faltantes="Protocolo operatorio.",
        recomendacion_interna="Revisión jurídica y técnica recomendada.",
    )


def test_informe_interno_contiene_la_frase_central_mandatada(tmp_path):
    item = _item()
    ruta = tmp_path / "informe_interno.docx"
    generar_informe_interno_docx(_caso_falso(), [item], [_hallazgo(item)], [], str(ruta))

    documento = DocxDocument(str(ruta))
    texto = "\n".join(p.text for p in documento.paragraphs)

    assert "hallazgo potencialmente discutible" in texto.lower()


def test_informe_interno_contiene_el_detalle_completo(tmp_path):
    item = _item()
    ruta = tmp_path / "informe_interno.docx"
    generar_informe_interno_docx(_caso_falso(), [item], [_hallazgo(item)], [], str(ruta))

    documento = DocxDocument(str(ruta))
    texto_tablas = "\n".join(
        celda.text for tabla in documento.tables for fila in tabla.rows for celda in fila.cells
    )

    assert "220305" in texto_tablas
    assert "dispositivo_medico" in texto_tablas


def test_informe_interno_incluye_fundamento_normativo_general(tmp_path):
    item = _item()
    ruta = tmp_path / "informe_interno.docx"
    generar_informe_interno_docx(_caso_falso(), [item], [_hallazgo(item)], [], str(ruta))

    documento = DocxDocument(str(ruta))
    texto_tablas = "\n".join(
        celda.text for tabla in documento.tables for fila in tabla.rows for celda in fila.cells
    )

    # El hallazgo de prueba es "dispositivo_medico": su fundamento normativo
    # general debe aparecer en la matriz de hallazgos del informe interno.
    assert "Compendio de Beneficios" in texto_tablas
    assert "Garantías Explícitas en Salud" in texto_tablas
