"""Verifica que la propuesta comercial NUNCA filtre metodología, códigos,
fundamentos jurídicos específicos, jurisprudencia ni estrategia de
reclamación, conforme a la separación estricta exigida entre el informe
interno y la propuesta comercial.
"""

import re
import types

import pymupdf
from docx import Document as DocxDocument

from app.reports.commercial_proposal_docx import generar_propuesta_docx
from app.reports.commercial_proposal_pdf import generar_propuesta_pdf
from app.reports.texts import GASTOS_EXTERNOS, REGLAS_PROPUESTA_EXTERNA

TERMINOS_PROHIBIDOS = [
    "codigo_prestacion",
    "código de prestación",
    "artículo",
    "ley n",
    "corte suprema",
    "superintendencia de salud",
    "glosa genérica",
    "demanda arbitral",
    "recurso de reposición",
    "hallazgo potencialmente discutible",
    "stent coronario",
    "arancel",
    "caec",
    "compendio de beneficios",
    "dfl n",
]

# "ges" (programa GES) se verifica con límite de palabra: como substring colisiona
# con "gestión"/"gestionar", que sí son parte del lenguaje comercial legítimo de la propuesta.
PATRON_ACRONIMO_GES = re.compile(r"\bges\b")


def _caso_falso():
    return types.SimpleNamespace(id=1, nombre_cliente="Cliente de Prueba", numero_cuenta="ABC-123")


def _honorarios(**overrides):
    base = {
        "honorario_fijo_texto": "UF 15",
        "honorario_exito": True,
        "honorario_exito_texto": "15% del beneficio económico obtenido",
        "exclusiones_texto": "",
    }
    base.update(overrides)
    return base


def test_propuesta_docx_no_filtra_metodologia(tmp_path):
    ruta = tmp_path / "propuesta.docx"
    generar_propuesta_docx(_caso_falso(), _honorarios(), str(ruta))

    documento = DocxDocument(str(ruta))
    texto = "\n".join(p.text for p in documento.paragraphs).lower()

    for termino in TERMINOS_PROHIBIDOS:
        assert termino.lower() not in texto
    assert PATRON_ACRONIMO_GES.search(texto) is None


def test_propuesta_pdf_no_filtra_metodologia_y_respeta_dos_paginas(tmp_path):
    ruta = tmp_path / "propuesta.pdf"
    generar_propuesta_pdf(_caso_falso(), _honorarios(honorario_exito=False, honorario_exito_texto=""), str(ruta))

    documento_pdf = pymupdf.open(str(ruta))
    texto = "\n".join(pagina.get_text() for pagina in documento_pdf).lower()
    n_paginas = documento_pdf.page_count
    documento_pdf.close()

    for termino in TERMINOS_PROHIBIDOS:
        assert termino.lower() not in texto
    assert PATRON_ACRONIMO_GES.search(texto) is None
    assert n_paginas <= 2


def test_propuesta_incluye_frase_obligatoria_y_ausencia_de_garantia(tmp_path):
    ruta = tmp_path / "propuesta.docx"
    generar_propuesta_docx(_caso_falso(), _honorarios(), str(ruta))

    documento = DocxDocument(str(ruta))
    texto = "\n".join(p.text for p in documento.paragraphs)

    assert "Del examen preliminar de los antecedentes recibidos" in texto
    assert "no asegura un resultado determinado" in texto
    assert GASTOS_EXTERNOS in texto
    assert "El encargo se ejecutará con base en los antecedentes que proporcione el cliente" in texto


def test_reglas_propuesta_externa_definidas():
    assert len(REGLAS_PROPUESTA_EXTERNA) >= 5
    assert all(isinstance(regla, str) and regla.strip() for regla in REGLAS_PROPUESTA_EXTERNA)
