"""Generador de la propuesta comercial (DOCX) para el potencial cliente.

ADVERTENCIA DE DISEÑO: este módulo debe respetar siempre
``app.reports.texts.REGLAS_PROPUESTA_EXTERNA``. Solo debe recibir textos
genéricos de honorarios configurados por el abogado. Ver
``tests/test_commercial_proposal_no_leak.py``.
"""

import datetime as dt

import docx

from app.reports import texts


def generar_propuesta_docx(caso, honorarios: dict, ruta_salida: str, vigencia_dias: int = 15) -> str:
    documento = docx.Document()

    documento.add_heading("Propuesta de Servicios Profesionales", level=1)
    documento.add_paragraph(f"Fecha: {dt.date.today().strftime('%d-%m-%Y')}")
    documento.add_paragraph(f"Cliente: {caso.nombre_cliente}")
    if caso.numero_cuenta:
        documento.add_paragraph(f"Caso relacionado: Cuenta clínica N° {caso.numero_cuenta}")

    documento.add_heading("Antecedentes", level=2)
    documento.add_paragraph(texts.FRASE_PROPUESTA_COMERCIAL)

    documento.add_heading("Servicios incluidos", level=2)
    for servicio in texts.SERVICIOS_PROPUESTA:
        documento.add_paragraph(servicio, style="List Bullet")

    documento.add_heading("Honorarios", level=2)
    documento.add_paragraph(
        f"Honorario fijo: {honorarios.get('honorario_fijo_texto', 'A definir')} (neto, más IVA)."
    )
    if honorarios.get("honorario_exito"):
        documento.add_paragraph(f"Honorario de éxito: {honorarios.get('honorario_exito_texto', 'A definir')}")
        documento.add_paragraph(texts.DEFINICION_EXITO)
    if honorarios.get("gastos_texto"):
        documento.add_paragraph(f"Gastos: {honorarios.get('gastos_texto')}")

    documento.add_heading("Exclusiones", level=2)
    documento.add_paragraph(honorarios.get("exclusiones_texto") or texts.EXCLUSIONES_POR_DEFECTO)

    documento.add_heading("Condiciones", level=2)
    documento.add_paragraph(texts.AUSENCIA_GARANTIA)
    documento.add_paragraph(
        f"Esta propuesta tiene una vigencia de {vigencia_dias} días corridos desde su fecha de emisión."
    )

    documento.add_heading("Aceptación", level=2)
    documento.add_paragraph("Nombre: ______________________________")
    documento.add_paragraph("RUT: ______________________________")
    documento.add_paragraph("Firma: ______________________________")
    documento.add_paragraph("Fecha: ______________________________")

    documento.save(ruta_salida)
    return ruta_salida
