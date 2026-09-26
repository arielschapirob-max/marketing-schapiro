"""Generador del informe interno jurídico-técnico (DOCX) — CONFIDENCIAL."""

import datetime as dt

import docx

from app.analysis.findings_engine import calcular_monto_total_discutible
from app.analysis.legal_grounds import fundamento_normativo_de
from app.extraction.document_types import TIPOS_DOCUMENTO, TIPOS_RELEVANTES_PARA_CASO
from app.extraction.normalization import formatear_rut
from app.reports import texts


def _fmt(monto) -> str:
    if monto is None:
        return "-"
    return f"${monto:,.0f}".replace(",", ".")


def _confianza_resumen(item) -> str:
    confianza = item.get_confianza()
    if not confianza:
        return "N/D"
    niveles = list(confianza.values())
    if "bajo" in niveles:
        return "bajo"
    if "medio" in niveles:
        return "medio"
    return "alto"


def generar_informe_interno_docx(caso, items, hallazgos, historial, ruta_salida: str, documentos=None) -> str:
    documentos = documentos or []
    documento = docx.Document()
    documento.add_heading("Informe Interno Jurídico-Técnico — CONFIDENCIAL", level=1)
    documento.add_paragraph(f"Fecha de generación: {dt.datetime.now().strftime('%d-%m-%Y %H:%M')}")
    documento.add_paragraph(f"Caso: {caso.nombre_cliente} — Cuenta N° {caso.numero_cuenta or 'N/D'}")
    documento.add_paragraph(f"Isapre: {caso.isapre or 'N/D'}   Prestador: {caso.prestador or 'N/D'}")
    if caso.rut_cliente:
        documento.add_paragraph(f"RUT afiliado: {formatear_rut(caso.rut_cliente)}")

    documento.add_heading("1. Resumen de la cuenta", level=2)
    total_cobrado = sum((i.valor_cobrado or 0) for i in items)
    total_bonificado = sum((i.valor_bonificado or 0) for i in items)
    documento.add_paragraph(f"Total de ítems registrados: {len(items)}")
    documento.add_paragraph(f"Valor total cobrado: {_fmt(total_cobrado)}")
    documento.add_paragraph(f"Valor total bonificado: {_fmt(total_bonificado)}")

    documento.add_heading("2. Matriz de ítems", level=2)
    tabla_items = documento.add_table(rows=1, cols=11)
    tabla_items.style = "Light Grid Accent 1"
    encabezados = [
        "Código",
        "Descripción",
        "Cobrado",
        "Bonificado",
        "Copago",
        "No cubierto",
        "Glosa",
        "Documento origen",
        "Página",
        "Confianza",
        "Aprobado",
    ]
    for celda, texto in zip(tabla_items.rows[0].cells, encabezados):
        celda.text = texto
    for item in items:
        fila = tabla_items.add_row().cells
        fila[0].text = item.codigo_prestacion or "-"
        fila[1].text = (item.descripcion or "-")[:80]
        fila[2].text = _fmt(item.valor_cobrado)
        fila[3].text = _fmt(item.valor_bonificado)
        fila[4].text = _fmt(item.copago)
        fila[5].text = _fmt(item.monto_no_cubierto)
        fila[6].text = (item.glosa or "-")[:60]
        fila[7].text = item.documento.nombre_archivo if item.documento else "-"
        fila[8].text = str(item.pagina_origen) if item.pagina_origen else "-"
        fila[9].text = _confianza_resumen(item)
        fila[10].text = "Sí" if item.aprobado else "No"

    documento.add_heading("3. Matriz de hallazgos", level=2)
    documento.add_paragraph(
        "Cada fila de la siguiente matriz corresponde a un hallazgo potencialmente discutible "
        "identificado en la cuenta, sujeto a validación profesional del abogado responsable."
    )
    documento.add_paragraph(
        "El fundamento normativo indicado es general (ley, DFL o compendio de la Superintendencia de Salud "
        "aplicable a la materia) y no incluye jurisprudencia: la incorporación de fallos o dictámenes verificados "
        "queda a criterio del abogado responsable. Los compendios de la Superintendencia se actualizan con "
        "frecuencia — verificar la redacción vigente antes de citarlos en una gestión formal."
    )
    tabla_hallazgos = documento.add_table(rows=1, cols=7)
    tabla_hallazgos.style = "Light Grid Accent 1"
    encabezados_h = [
        "Tipo",
        "Ítem",
        "Monto discutible",
        "Prioridad",
        "Estado",
        "Explicación interna",
        "Fundamento normativo (general — verificar vigencia)",
    ]
    for celda, texto in zip(tabla_hallazgos.rows[0].cells, encabezados_h):
        celda.text = texto
    for h in hallazgos:
        fila = tabla_hallazgos.add_row().cells
        fila[0].text = h.tipo
        fila[1].text = str(h.item_id) if h.item_id else "General"
        fila[2].text = _fmt(h.monto_discutible)
        fila[3].text = h.prioridad
        fila[4].text = h.estado
        fila[5].text = h.explicacion_interna or ""
        fila[6].text = fundamento_normativo_de(h.tipo)

    documento.add_heading("4. Monto potencialmente discutible (sin duplicar ítems)", level=2)
    monto_total = calcular_monto_total_discutible(hallazgos)
    documento.add_paragraph(f"Monto potencialmente discutible estimado: {_fmt(monto_total)}")

    documento.add_heading("5. Antecedentes disponibles y faltantes", level=2)
    tipos_disponibles = {d.tipo_documento for d in documentos}
    documento.add_paragraph("Documentos del caso:")
    for tipo_clave in TIPOS_RELEVANTES_PARA_CASO:
        disponible = tipo_clave in tipos_disponibles
        etiqueta = TIPOS_DOCUMENTO.get(tipo_clave, tipo_clave)
        documento.add_paragraph(f"{'Disponible' if disponible else 'FALTANTE'} — {etiqueta}", style="List Bullet")
    faltantes_por_hallazgo = sorted({h.documentos_faltantes for h in hallazgos if h.documentos_faltantes})
    if faltantes_por_hallazgo:
        documento.add_paragraph("Antecedentes adicionales solicitados por hallazgos específicos:")
        for f in faltantes_por_hallazgo:
            documento.add_paragraph(f, style="List Bullet")

    documento.add_heading("6. Advertencia de validación profesional", level=2)
    parrafo = documento.add_paragraph(texts.ADVERTENCIA_VALIDACION_PROFESIONAL)
    parrafo.runs[0].bold = True

    documento.add_heading("7. Historial de modificaciones y aprobación", level=2)
    if historial:
        for cambio in historial:
            documento.add_paragraph(
                f"[{cambio.fecha.strftime('%d-%m-%Y %H:%M')}] {cambio.usuario} — {cambio.entidad} "
                f"#{cambio.entidad_id} · {cambio.campo}: '{cambio.valor_anterior}' → '{cambio.valor_nuevo}'"
            )
    else:
        documento.add_paragraph("Sin modificaciones registradas.")

    aprobacion = f"Aprobado por el abogado responsable: {'Sí' if caso.aprobado_por_abogado else 'No'}"
    if caso.fecha_aprobacion:
        aprobacion += f" — {caso.fecha_aprobacion.strftime('%d-%m-%Y %H:%M')}"
    documento.add_paragraph(aprobacion)

    documento.save(ruta_salida)
    return ruta_salida
