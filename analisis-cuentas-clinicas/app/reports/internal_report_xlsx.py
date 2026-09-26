"""Generador del informe interno jurídico-técnico (XLSX) — CONFIDENCIAL."""

import openpyxl
from openpyxl.styles import Font

from app.analysis.findings_engine import calcular_monto_total_discutible
from app.analysis.legal_grounds import fundamento_normativo_de
from app.extraction.document_types import TIPOS_DOCUMENTO, TIPOS_RELEVANTES_PARA_CASO


def _confianza_resumen(item) -> str:
    confianza = item.get_confianza()
    if not confianza:
        return "N/D"
    niveles = list(confianza.values())
    if "bajo" in niveles:
        return "bajo"
    if "medio" in niveles:
        return "medio"
    return "alto"


def generar_informe_interno_xlsx(caso, items, hallazgos, historial, ruta_salida: str, documentos=None) -> str:
    documentos = documentos or []
    wb = openpyxl.Workbook()

    resumen = wb.active
    resumen.title = "Resumen"
    resumen.append(["Informe Interno Jurídico-Técnico — CONFIDENCIAL"])
    resumen["A1"].font = Font(bold=True, size=14)
    resumen.append(["Cliente", caso.nombre_cliente])
    resumen.append(["Cuenta", caso.numero_cuenta or "N/D"])
    resumen.append(["Isapre", caso.isapre or "N/D"])
    resumen.append(["Prestador", caso.prestador or "N/D"])
    total_cobrado = sum((i.valor_cobrado or 0) for i in items)
    total_bonificado = sum((i.valor_bonificado or 0) for i in items)
    monto_discutible = calcular_monto_total_discutible(hallazgos)
    resumen.append(["Total cobrado", total_cobrado])
    resumen.append(["Total bonificado", total_bonificado])
    resumen.append(["Monto potencialmente discutible (sin duplicar ítems)", monto_discutible])

    resumen.append([])
    resumen.append(["Documentos del caso"])
    tipos_disponibles = {d.tipo_documento for d in documentos}
    for tipo_clave in TIPOS_RELEVANTES_PARA_CASO:
        etiqueta = TIPOS_DOCUMENTO.get(tipo_clave, tipo_clave)
        resumen.append([etiqueta, "Disponible" if tipo_clave in tipos_disponibles else "FALTANTE"])

    hoja_items = wb.create_sheet("Items")
    hoja_items.append(
        [
            "ID",
            "Código",
            "Descripción",
            "Cantidad",
            "Cobrado",
            "Bonificado",
            "Copago",
            "No cubierto",
            "Deducible",
            "Total",
            "Glosa",
            "Documento origen",
            "Página",
            "Confianza",
            "Aprobado",
            "Editado manualmente",
        ]
    )
    for item in items:
        hoja_items.append(
            [
                item.id,
                item.codigo_prestacion,
                item.descripcion,
                item.cantidad,
                item.valor_cobrado,
                item.valor_bonificado,
                item.copago,
                item.monto_no_cubierto,
                item.deducible,
                item.total,
                item.glosa,
                item.documento.nombre_archivo if item.documento else None,
                item.pagina_origen,
                _confianza_resumen(item),
                "Sí" if item.aprobado else "No",
                "Sí" if item.editado_manualmente else "No",
            ]
        )

    hoja_hallazgos = wb.create_sheet("Hallazgos")
    hoja_hallazgos.append(
        [
            "ID",
            "Tipo",
            "Item ID",
            "Monto discutible",
            "Prioridad",
            "Estado",
            "Explicación interna",
            "Documentos faltantes",
            "Recomendación interna",
            "Fundamento normativo (general — verificar vigencia)",
        ]
    )
    for h in hallazgos:
        hoja_hallazgos.append(
            [
                h.id,
                h.tipo,
                h.item_id,
                h.monto_discutible,
                h.prioridad,
                h.estado,
                h.explicacion_interna,
                h.documentos_faltantes,
                h.recomendacion_interna,
                fundamento_normativo_de(h.tipo),
            ]
        )

    hoja_historial = wb.create_sheet("Historial")
    hoja_historial.append(["Fecha", "Usuario", "Entidad", "Entidad ID", "Campo", "Valor anterior", "Valor nuevo"])
    for cambio in historial:
        hoja_historial.append(
            [
                cambio.fecha.strftime("%d-%m-%Y %H:%M"),
                cambio.usuario,
                cambio.entidad,
                cambio.entidad_id,
                cambio.campo,
                cambio.valor_anterior,
                cambio.valor_nuevo,
            ]
        )

    wb.save(ruta_salida)
    return ruta_salida
