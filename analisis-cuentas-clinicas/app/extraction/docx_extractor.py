"""Extracción de texto y tablas desde documentos DOCX, usando python-docx."""

import docx


def extraer_texto_docx(ruta: str) -> tuple[str, list]:
    documento = docx.Document(ruta)
    parrafos = [p.text for p in documento.paragraphs if p.text.strip()]
    tablas = []
    for tabla in documento.tables:
        filas = [[celda.text.strip() for celda in fila.cells] for fila in tabla.rows]
        tablas.append(filas)
    return "\n".join(parrafos), tablas
