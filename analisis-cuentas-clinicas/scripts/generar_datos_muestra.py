"""Genera documentos ficticios de prueba para analisis-cuentas-clinicas.

Los datos aquí contenidos son inventados y no corresponden a personas,
cuentas ni isapres reales. Uso exclusivo para pruebas locales de la
aplicación (extracción, revisión manual, hallazgos e informes).

Ejecutar desde la raíz del proyecto:
    python scripts/generar_datos_muestra.py
"""

from pathlib import Path

import docx
import openpyxl

DESTINO = Path(__file__).resolve().parent.parent / "samples"
DESTINO.mkdir(parents=True, exist_ok=True)


def generar_xlsx_muestra() -> Path:
    wb = openpyxl.Workbook()
    hoja = wb.active
    hoja.title = "Liquidacion"
    hoja.append(
        [
            "Código",
            "Descripción",
            "Cantidad",
            "Valor cobrado",
            "Valor bonificado",
            "Copago",
            "Glosa",
        ]
    )
    hoja.append(
        [
            "110101",
            "Consulta médica especialidad",
            1,
            45000,
            30000,
            15000,
            "Bonificación según plan",
        ]
    )
    hoja.append(["220305", "Stent coronario liberador de fármaco", 1, 3200000, 0, 3200000, "No corresponde"])
    hoja.append(["330501", "Día cama UCI adulto", 3, 1800000, 1200000, 600000, "Bonificación parcial"])
    hoja.append(["440102", "Examen de laboratorio general", 2, 25000, 25000, 0, "Bonificación 100%"])
    hoja.append(["440102", "Examen de laboratorio general", 2, 25000, 25000, 0, "Bonificación 100%"])
    ruta = DESTINO / "liquidacion_isapre_ficticia.xlsx"
    wb.save(ruta)
    return ruta


def generar_docx_muestra() -> Path:
    documento = docx.Document()
    documento.add_heading("Cuenta Clínica (documento ficticio de prueba)", level=1)
    documento.add_paragraph("Afiliado: Paciente de Prueba Ficticio")
    documento.add_paragraph("RUT: 11.111.111-1")
    documento.add_paragraph("Isapre: Isapre Ejemplo")
    documento.add_paragraph("Prestador: Clínica Ejemplo S.A.")
    documento.add_paragraph("Número de cuenta: CTA-2026-0001")
    documento.add_paragraph("Diagnóstico: Diagnóstico ficticio de prueba")

    tabla = documento.add_table(rows=1, cols=5)
    encabezados = ["Código", "Descripción", "Valor cobrado", "Valor bonificado", "Copago"]
    for celda, texto in zip(tabla.rows[0].cells, encabezados):
        celda.text = texto
    filas_muestra = [
        ["220305", "Stent coronario liberador de fármaco", "3.200.000", "0", "3.200.000"],
        ["330501", "Día cama UCI adulto", "1.800.000", "1.200.000", "600.000"],
    ]
    for fila_datos in filas_muestra:
        fila = tabla.add_row().cells
        for celda, valor in zip(fila, fila_datos):
            celda.text = valor

    ruta = DESTINO / "cuenta_clinica_ficticia.docx"
    documento.save(ruta)
    return ruta


if __name__ == "__main__":
    ruta_xlsx = generar_xlsx_muestra()
    ruta_docx = generar_docx_muestra()
    print(f"Generado: {ruta_xlsx}")
    print(f"Generado: {ruta_docx}")
    print("\nRecuerde: estos son datos ficticios de prueba. Nunca cargue datos reales en el repositorio.")
